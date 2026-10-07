from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


class ExcelExporter:
    def __init__(self):
        self.workbook = None
        self.header_font = Font(bold=True, color="FFFFFF")
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.cell_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

    def create_workbook(self):
        self.workbook = openpyxl.Workbook()

    def save(self, file_path: str):
        path = Path(file_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.workbook.save(path)

    def write_header_row(self, worksheet, headers: list[str], row: int = 1):
        for col_idx, header in enumerate(headers, 1):
            cell = worksheet.cell(row=row, column=col_idx, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.thin_border

    def write_data_row(
        self, worksheet, data: list, row: int, formats: dict[int, str] | None = None
    ):
        for col_idx, value in enumerate(data, 1):
            cell = worksheet.cell(row=row, column=col_idx, value=value)
            cell.alignment = self.cell_alignment
            cell.border = self.thin_border
            if formats and col_idx in formats:
                cell.number_format = formats[col_idx]

    def auto_fit_columns(self, worksheet, min_width: int = 10, max_width: int = 50):
        for col in worksheet.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                max_length = max(max_length, len(str(cell.value)))
            adjusted_width = min(max(max_length + 2, min_width), max_width)
            worksheet.column_dimensions[column].width = adjusted_width
